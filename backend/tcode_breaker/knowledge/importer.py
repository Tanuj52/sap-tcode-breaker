"""Universal knowledge importer: whatever the user has, turned into tables / function modules.

Accepted (auto-detected, also inside ``` fences copied from an AI chat):
  * delimited text (CSV, TSV, ';' or '|' separated) and Excel (.xlsx) with a header row:
      - SE16 export of DD03L (TABNAME, FIELDNAME, KEYFLAG, DATATYPE, LENG, DECIMALS, REFFIELD, DDTEXT ...)
      - SE11 field list copied for one table (Field, Key, Data element, Data Type, Length, Decimal Places, Short Description)
      - INFORMATION_SCHEMA.COLUMNS style (TABLE_NAME, COLUMN_NAME, DATA_TYPE, ...)
      - any similar header; see _ALIASES
  * markdown tables (| TABNAME | FIELDNAME | ... |) as returned by chat assistants
  * JSON / YAML: the pack format ({tables: {NAME: {fields: {F: "CHAR(10) Text"}}}}), a list of table objects
    ([{name, description, key, fields: [{name, type, length, decimals, key, description}]}]),
    or flat row objects; function modules as {function_modules: {NAME: {kind, output, sql, note}}}
  * ABAP source with FUNCTION ... ENDFUNCTION blocks (custom function modules)

Nothing is saved here: parse() returns a draft that the UI previews before it is written to the library.
"""

from __future__ import annotations

import csv
import io
import json
import re
from dataclasses import dataclass, field

import yaml

from .models import FM_KINDS, FieldInfo, FMInfo, Provenance, S4Mapping, TableInfo, fm_from_spec, parse_field_spec

_NAME = re.compile(r"^[A-Z0-9_/]{1,40}$")

# canonical column -> accepted header spellings (compared upper-case, spaces/dots/dashes removed)
_ALIASES = {
    "TABNAME": ["TABNAME", "TABLE", "TABLENAME", "TABLE_NAME", "SAPTABLE", "SAP_TABLE", "TAB", "OBJECT", "ENTITY"],
    "FIELDNAME": ["FIELDNAME", "FIELD", "FIELD_NAME", "COLUMN", "COLUMNNAME", "COLUMN_NAME", "FNAME", "COMPONENT", "NAME"],
    "KEYFLAG": ["KEYFLAG", "KEY", "ISKEY", "IS_KEY", "PRIMARYKEY", "PRIMARY_KEY", "PK", "KEYFIELD", "KEY_FIELD"],
    "DATATYPE": ["DATATYPE", "DATA_TYPE", "TYPE", "DDICTYPE", "DDIC_TYPE", "ABAPTYPE", "ABAP_TYPE", "INTTYPE", "SNOWFLAKE_TYPE"],
    "LENG": ["LENG", "LENGTH", "LEN", "SIZE", "CHARACTER_MAXIMUM_LENGTH", "NUMERIC_PRECISION", "PRECISION", "INTLEN", "NO_OF_CHARACTERS"],
    "DECIMALS": ["DECIMALS", "DECIMAL", "DECIMALPLACES", "DECIMAL_PLACES", "DEC", "SCALE", "NUMERIC_SCALE"],
    "REFFIELD": ["REFFIELD", "REF_FIELD", "REFERENCEFIELD", "REFERENCE_FIELD", "CURRENCY_UNIT_FIELD", "UNIT_FIELD"],
    "DDTEXT": ["DDTEXT", "DESCRIPTION", "FIELDDESCRIPTION", "FIELD_DESCRIPTION", "SHORTDESCRIPTION", "SHORT_DESCRIPTION", "TEXT",
               "SCRTEXT_M", "SCRTEXT_L", "LABEL", "MEANING", "COMMENT", "FIELD_TEXT", "FIELDTEXT", "SHORTTEXT"],
    "TABTEXT": ["TABTEXT", "TABLE_TEXT", "TABLETEXT", "TABLE_DESCRIPTION", "TABLEDESCRIPTION", "TABLE_COMMENT", "TABDESC"],
    "MODULE": ["MODULE", "APPLICATION", "APP_COMPONENT", "AREA"],
    "ROLLNAME": ["ROLLNAME", "DATAELEMENT", "DATA_ELEMENT"],
    "POSITION": ["POSITION", "POS", "ORDINAL_POSITION"],
}
_LOOKUP = {a.replace("_", ""): canon for canon, al in _ALIASES.items() for a in al}

# ABAP / DDIC / Snowflake type spellings -> DDIC data type used by the knowledge base
_TYPE_ALIASES = {
    "C": "CHAR", "CHARACTER": "CHAR", "N": "NUMC", "D": "DATS", "DATE": "DATS", "DATN": "DATS", "T": "TIMS", "TIME": "TIMS", "TIMN": "TIMS",
    "P": "DEC", "PACKED": "DEC", "NUMERIC": "DEC", "DECIMAL": "DEC", "NUMBER": "DEC", "DF16_DEC": "DEC", "DF34_DEC": "DEC",
    "DF16_RAW": "DEC", "DF34_RAW": "DEC", "D34N": "D16N", "DECFLOAT16": "D16N", "DECFLOAT34": "D16N",
    "I": "INT4", "INT": "INT4", "INTEGER": "INT4", "B": "INT1", "S": "INT2", "BIGINT": "INT8", "SMALLINT": "INT2", "TINYINT": "INT1",
    "F": "FLTP", "FLOAT": "FLTP", "DOUBLE": "FLTP", "REAL": "FLTP", "X": "RAW", "LRAW": "RAW", "RAWSTRING": "RAW", "XSTRING": "RAW",
    "BINARY": "RAW", "VARBINARY": "RAW", "STRING": "STRG", "RSTR": "STRG", "TEXT": "STRG", "LCHR": "CHAR", "ACCP": "NUMC", "PREC": "INT2",
    "VARCHAR": "CHAR", "NVARCHAR": "CHAR", "NCHAR": "CHAR", "BOOLEAN": "CHAR", "BOOL": "CHAR", "UTCLONG": "TIMESTAMP", "UTCL": "TIMESTAMP",
    "TIMESTAMP_NTZ": "TIMESTAMP", "TIMESTAMP_LTZ": "TIMESTAMP", "TIMESTAMP_TZ": "TIMESTAMP", "DATETIME": "TIMESTAMP", "CURRENCY": "CURR",
    "QUANTITY": "QUAN", "LANGUAGE": "LANG", "CLIENT": "CLNT",
}
_TYPE_SPEC = re.compile(r"^\s*([A-Za-z_0-9]+)\s*(?:[(\s]\s*(\d+)\s*(?:[,.\s]\s*(\d+))?\s*\)?)?\s*$")
_TRUE = {"X", "TRUE", "YES", "Y", "1", "KEY", "PK", "✓", "✔", "☑", "T"}


@dataclass
class ImportDraft:
    format: str = "unknown"
    tables: dict[str, TableInfo] = field(default_factory=dict)
    fms: dict[str, FMInfo] = field(default_factory=dict)
    methods: dict[str, FMInfo] = field(default_factory=dict)
    fm_sources: dict[str, str] = field(default_factory=dict)
    s4: list[S4Mapping] = field(default_factory=list)
    unknown: list[str] = field(default_factory=list)  # names an AI said it does not know
    warnings: list[str] = field(default_factory=list)

    def merge(self, other: "ImportDraft") -> None:
        for name, t in other.tables.items():
            if name in self.tables:
                self.tables[name].fields.update(t.fields)
                self.tables[name].key = self.tables[name].key or t.key
                self.tables[name].description = self.tables[name].description or t.description
            else:
                self.tables[name] = t
        self.fms.update(other.fms)
        self.methods.update(other.methods)
        self.fm_sources.update(other.fm_sources)
        self.s4 += other.s4
        self.unknown += other.unknown
        self.warnings += other.warnings
        self.format = other.format if self.format == "unknown" else (self.format if other.format in (self.format, "unknown") else "mixed")

    @property
    def empty(self) -> bool:
        return not (self.tables or self.fms or self.methods or self.fm_sources or self.s4)


# ---------------------------------------------------------------------- entry points
def parse(data: str | bytes, filename: str | None = None, default_table: str | None = None, origin: str = "import") -> ImportDraft:
    """Parse any supported input. `default_table` names the table when rows carry no table column (SE11 copy)."""
    name = (filename or "").lower()
    if isinstance(data, bytes):
        if data[:2] == b"PK":
            if name.endswith((".xlsx", ".xlsm")) or (not name.endswith(".zip") and b"[Content_Types].xml" in data[:4000]):
                return _finish(_parse_rows(_xlsx_rows(data), default_table, "excel"), origin)
            return _finish(_parse_zip(data, default_table, origin), origin)
        data = data.decode("utf-8-sig", errors="replace")
    text = data.replace("\r\n", "\n").replace("\r", "\n").lstrip("\ufeff")
    if name.endswith(".abap") or re.search(r"^\s*FUNCTION\s+[A-Za-z0-9_/]+\s*\.", text, re.IGNORECASE | re.MULTILINE):
        from .library import split_function_modules

        blocks = split_function_modules(text)
        if blocks:
            return _finish(ImportDraft(format="abap", fm_sources=blocks), origin)
    fenced = re.findall(r"```[A-Za-z0-9_-]*\n(.*?)```", text, re.DOTALL)
    if fenced:
        draft = ImportDraft()
        for block in fenced:
            draft.merge(_parse_text(block, default_table))
        return _finish(draft, origin)
    return _finish(_parse_text(text, default_table), origin)


def _parse_text(text: str, default_table: str | None) -> ImportDraft:
    s = text.strip()
    if not s:
        return ImportDraft(warnings=["Nothing to import"])
    if s[0] in "[{":
        try:
            return _parse_structured(json.loads(s), "json", default_table)
        except json.JSONDecodeError as exc:
            if s[0] == "{" or s.rstrip()[-1] in "]}":
                # truncated / sloppy JSON from a chat: try YAML (superset) before giving up
                try:
                    return _parse_structured(yaml.safe_load(s), "json", default_table)
                except yaml.YAMLError:
                    return ImportDraft(warnings=[f"Looks like JSON but does not parse: {exc}"])
    lines = [ln for ln in s.split("\n") if ln.strip()]
    pipe_rows = [ln for ln in lines if ln.strip().startswith("|")]
    if len(pipe_rows) >= 2 and len(pipe_rows) >= len(lines) * 0.6:
        return _parse_rows(_markdown_rows(pipe_rows), default_table, "markdown")
    if re.search(r"^(tables|function_modules|methods|mappings)\s*:", s, re.MULTILINE):
        try:
            return _parse_structured(yaml.safe_load(s), "yaml", default_table)
        except yaml.YAMLError as exc:
            return ImportDraft(warnings=[f"YAML syntax error: {exc}"])
    rows = _delimited_rows(s)
    if rows and len(rows[0]) > 1:
        return _parse_rows(rows, default_table, "delimited")
    try:
        doc = yaml.safe_load(s)
        if isinstance(doc, (dict, list)):
            return _parse_structured(doc, "yaml", default_table)
    except yaml.YAMLError:
        pass
    return ImportDraft(warnings=["Could not recognise the format. Use a header row (TABNAME, FIELDNAME, DATATYPE, LENG, DECIMALS, KEYFLAG, DDTEXT), "
                                 "JSON/YAML in the pack format, or FUNCTION ... ENDFUNCTION source."])


def _finish(d: ImportDraft, origin: str) -> ImportDraft:
    for t in d.tables.values():
        t.origin = origin
        if origin == "ai":
            t.source, t.fields_complete = Provenance.AI, False
            for f in t.fields.values():
                f.source = Provenance.AI
        if not t.key:
            t.key = [n for n, f in t.fields.items() if f.is_key]
        for k in t.key:
            if k in t.fields:
                t.fields[k].is_key = True
        if not t.fields:
            d.warnings.append(f"{t.name}: no fields")
        elif not t.key:
            d.warnings.append(f"{t.name}: no key fields marked (joins and de-duplication are less certain)")
    for f in list(d.fms.values()) + list(d.methods.values()):
        f.origin = origin
        if origin == "ai":
            f.source = Provenance.AI
    if d.empty and not d.unknown and not d.warnings:
        d.warnings.append("Nothing recognisable found")
    return d


# ---------------------------------------------------------------------- tabular
def _delimited_rows(text: str) -> list[list[str]]:
    sample = "\n".join(text.split("\n")[:20])
    counts = {sep: min((ln.count(sep) for ln in sample.split("\n") if ln.strip()), default=0) for sep in ("\t", ";", ",", "|")}
    sep = max(counts, key=lambda k: (counts[k], k == "\t"))
    if counts[sep] == 0:
        return []
    return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), delimiter=sep) if any(c.strip() for c in row)]


def _markdown_rows(lines: list[str]) -> list[list[str]]:
    rows = []
    for ln in lines:
        cells = [c.strip().strip("`*") for c in ln.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{2,}:?", c) or not c for c in cells):
            continue
        rows.append(cells)
    return rows


def _xlsx_rows(data: bytes) -> list[list[str]]:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - dependency is declared, guard anyway
        raise ValueError("Excel import needs the openpyxl package; save the sheet as CSV instead") from exc
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    rows: list[list[str]] = []
    for ws in wb.worksheets:
        for r in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v).strip() for v in r]
            if any(cells):
                rows.append(cells)
        if rows:
            break
    return rows


def _header_map(row: list[str]) -> dict[int, str]:
    out = {}
    for i, h in enumerate(row):
        k = re.sub(r"[\s.\-_/]+", "", (h or "").upper())
        canon = _LOOKUP.get(k)
        if canon and canon not in out.values():
            out[i] = canon
    return out


def _parse_rows(rows: list[list[str]], default_table: str | None, fmt: str) -> ImportDraft:
    d = ImportDraft(format=fmt)
    header_at, hmap = None, {}
    for i, row in enumerate(rows[:15]):
        m = _header_map(row)
        if "FIELDNAME" in m.values() and ("DATATYPE" in m.values() or "TABNAME" in m.values() or "DDTEXT" in m.values()):
            header_at, hmap = i, m
            break
    if header_at is None:
        d.warnings.append("No header row found. The first row must name the columns, e.g. TABNAME, FIELDNAME, DATATYPE, LENG, DECIMALS, KEYFLAG, DDTEXT")
        return d
    if "TABNAME" not in hmap.values() and not default_table:
        d.warnings.append("Rows have no table column: enter the table name these fields belong to")
        return d
    width = len(rows[header_at])
    for n, row in enumerate(rows[header_at + 1:], start=header_at + 2):
        row = _realign(row, width, hmap)
        rec = {canon: (row[i] if i < len(row) else "").strip() for i, canon in hmap.items()}
        if not any(rec.values()) or _looks_like_header(rec):
            continue
        tab = (rec.get("TABNAME") or default_table or "").upper()
        fname = (rec.get("FIELDNAME") or "").upper()
        if fname.startswith((".INCLU", ".APPEND", ".INCLUDE")) or not fname:
            continue
        if not _NAME.match(tab) or not _NAME.match(fname):
            d.warnings.append(f"Row {n}: skipped invalid name {tab}.{fname}")
            continue
        t = d.tables.setdefault(tab, TableInfo(name=tab, source=Provenance.LIBRARY, fields_complete=False))
        if rec.get("TABTEXT") and not t.description:
            t.description = rec["TABTEXT"]
        if rec.get("MODULE") and not t.module:
            t.module = rec["MODULE"].upper()
        ddic, length, dec, warn = normalize_type(rec.get("DATATYPE"), rec.get("LENG"), rec.get("DECIMALS"))
        if warn:
            d.warnings.append(f"{tab}.{fname}: {warn}")
        is_key = (rec.get("KEYFLAG") or "").strip().upper() in _TRUE
        if fname in t.fields:
            d.warnings.append(f"{tab}.{fname}: listed twice, last row kept")
        t.fields[fname] = FieldInfo(name=fname, ddic_type=ddic, length=length, decimals=dec, ref_field=(rec.get("REFFIELD") or "").upper() or None,
                                    description=rec.get("DDTEXT") or "", is_key=is_key, source=Provenance.LIBRARY)
        if is_key and fname not in t.key:
            t.key.append(fname)
    return d


def _realign(row: list[str], width: int, hmap: dict[int, str]) -> list[str]:
    """Rows with stray empty cells (common in AI output: `T,F,,,CHAR,20`) are shifted back so the
    data type lands in the DATATYPE column. Only empty cells are ever removed."""
    extra = len(row) - width
    ti = next((i for i, c in hmap.items() if c == "DATATYPE"), None)
    if extra <= 0 or ti is None:
        return row
    from .kb import BASE_DIR

    valid = _valid_types(BASE_DIR)
    if (row[ti] if ti < len(row) else "").upper() in valid:
        return row
    for j in range(ti, min(len(row), ti + extra + 1)):
        if row[j].upper() in valid and all(not c for c in row[ti:j]):
            return row[:ti] + row[j:]
    return row


def _looks_like_header(rec: dict) -> bool:
    return (rec.get("FIELDNAME") or "").upper() in ("FIELDNAME", "FIELD", "COLUMN_NAME") and (rec.get("TABNAME") or "").upper() in ("", "TABNAME", "TABLE")


def normalize_type(raw: str | None, length: str | int | None = None, decimals: str | int | None = None) -> tuple[str, int | None, int | None, str | None]:
    """'CHAR(10)', 'char 10', 'C', 'NUMBER(13,2)', 'VARCHAR(40)', 'DATS' ... -> (DDIC type, length, decimals, warning)."""
    from .kb import BASE_DIR

    valid = _valid_types(BASE_DIR)
    ln, dc = _int(length), _int(decimals)
    s = (raw or "").strip().upper()
    if not s:
        return "CHAR", ln, dc, "no data type given, CHAR assumed"
    m = _TYPE_SPEC.match(s)
    if m:
        s = m.group(1)
        ln = ln if ln is not None else _int(m.group(2))
        dc = dc if dc is not None else _int(m.group(3))
    t = _TYPE_ALIASES.get(s, s)
    if t == "DEC" and s in ("NUMBER", "NUMERIC", "DECIMAL") and not dc:
        t = "INT8" if (ln or 0) > 10 else "INT4"
    if t not in valid:
        return "CHAR", ln, dc, f"type {raw!r} not recognised, CHAR assumed"
    if t in ("CHAR", "NUMC") and ln is None:
        return t, None, None, "no length given"
    return t, ln, dc if dc else None, None


_VALID: set[str] | None = None


def _valid_types(base) -> set[str]:
    global _VALID
    if _VALID is None:
        _VALID = set((yaml.safe_load((base / "ddic_type_map.yaml").read_text()) or {}).get("types", {}))
    return _VALID


def _int(v) -> int | None:
    try:
        return int(float(str(v).strip())) if v not in (None, "") and str(v).strip() else None
    except ValueError:
        return None


# ---------------------------------------------------------------------- structured (JSON / YAML)
def _parse_structured(doc, fmt: str, default_table: str | None) -> ImportDraft:
    d = ImportDraft(format=fmt)
    if isinstance(doc, list):
        doc = {"tables": doc} if doc and isinstance(doc[0], dict) and not _is_fm(doc[0]) else {"function_modules": doc}
    if not isinstance(doc, dict):
        d.warnings.append("Expected an object or a list")
        return d
    keys = {k.lower(): k for k in doc}
    found = False
    if "tables" in keys:
        found = True
        _tables_from(doc[keys["tables"]], d)
    for sec in ("function_modules", "fms", "functions", "methods"):
        if sec in keys:
            found = True
            target = d.methods if sec == "methods" else d.fms
            _fms_from(doc[keys[sec]], d, target)
    for sec in ("mappings", "s4_mappings"):
        if sec in keys:
            found = True
            for i, m in enumerate(doc[keys[sec]] or []):
                try:
                    d.s4.append(S4Mapping(**m))
                except Exception as exc:  # noqa: BLE001
                    d.warnings.append(f"S/4 mapping #{i + 1}: {exc}")
    if "unknown" in keys:
        d.unknown += [str(x).upper() for x in doc[keys["unknown"]] or []]
    if not found:
        if all(isinstance(v, dict) for v in doc.values()) and doc:
            if all(_is_fm(v) for v in doc.values()):
                _fms_from(doc, d, d.fms)
            else:
                _tables_from(doc, d)
        elif ("name" in keys or "table" in keys) and ("fields" in keys or "columns" in keys):
            _tables_from([doc], d)
        elif default_table and ("fields" in keys or "columns" in keys):
            _tables_from([{"name": default_table, **doc}], d)
        else:
            d.warnings.append("No 'tables' or 'function_modules' section found")
    return d


def _is_fm(obj: dict) -> bool:
    return isinstance(obj, dict) and "kind" in {k.lower() for k in obj} and "fields" not in {k.lower() for k in obj}


def _get(obj: dict, *names, default=None):
    low = {re.sub(r"[\s_]+", "", k.lower()): v for k, v in obj.items()}
    for n in names:
        v = low.get(n.replace("_", ""))
        if v not in (None, ""):
            return v
    return default


def _tables_from(section, d: ImportDraft) -> None:
    items = section.items() if isinstance(section, dict) else [(None, x) for x in section or []]
    for name, spec in items:
        if not isinstance(spec, dict):
            d.warnings.append(f"{name}: entry is not an object")
            continue
        tname = str(name or _get(spec, "name", "table", "tabname", "table_name", default="")).upper().strip()
        if _get(spec, "unknown") in (True, "true", "yes"):
            d.unknown.append(tname)
            continue
        if not _NAME.match(tname):
            d.warnings.append(f"Skipped table with invalid name {tname!r}")
            continue
        key_raw = _get(spec, "key", "keys", "primary_key", "primarykey", default=[])
        key = [k.strip().upper() for k in (key_raw.split(",") if isinstance(key_raw, str) else key_raw or [])]
        t = TableInfo(name=tname, description=str(_get(spec, "description", "text", "ddtext", "title", "comment", default="")),
                      module=str(_get(spec, "module", "application", default="")).upper(), key=key,
                      release=str(_get(spec, "release", default="both")).lower() if str(_get(spec, "release", default="both")).lower() in ("ecc", "s4", "both") else "both",
                      fields_complete=_get(spec, "fields_complete", default=False) in (True, "true", "yes", "X"),
                      text_table=_get(spec, "text_table"), text_of=_get(spec, "text_of"), source=Provenance.LIBRARY)
        conf = str(_get(spec, "confidence", default="")).lower()
        if conf in ("medium", "low"):
            d.warnings.append(f"{tname}: AI confidence {conf} - verify against SE11 before relying on it")
        fields = _get(spec, "fields", "columns", default={})
        fitems = fields.items() if isinstance(fields, dict) else [(None, x) for x in fields or []]
        for fname, fspec in fitems:
            f = _field_from(fname, fspec, key, d, tname)
            if f is not None:
                t.fields[f.name] = f
        d.tables[tname] = t


def _field_from(fname, fspec, key: list[str], d: ImportDraft, tname: str) -> FieldInfo | None:
    if isinstance(fspec, str) or fspec is None:
        name = str(fname or "").upper()
        if not _NAME.match(name):
            d.warnings.append(f"{tname}: skipped field {name!r}")
            return None
        f = parse_field_spec(name, fspec or "", Provenance.LIBRARY, key)
        ddic, ln, dc, warn = normalize_type(f.ddic_type, f.length, f.decimals)
        if warn and warn != "no length given":
            d.warnings.append(f"{tname}.{name}: {warn}")
        return f.model_copy(update={"ddic_type": ddic, "length": ln, "decimals": dc})
    if not isinstance(fspec, dict):
        return None
    name = str(fname or _get(fspec, "name", "field", "fieldname", "column", "column_name", default="")).upper().strip()
    if not _NAME.match(name):
        d.warnings.append(f"{tname}: skipped field {name!r}")
        return None
    ddic, ln, dc, warn = normalize_type(str(_get(fspec, "type", "datatype", "data_type", "ddic_type", "abap_type", default="") or ""),
                                        _get(fspec, "length", "leng", "len", "size"), _get(fspec, "decimals", "dec", "scale"))
    if warn and warn != "no length given":
        d.warnings.append(f"{tname}.{name}: {warn}")
    k = _get(fspec, "key", "is_key", "keyflag", "primary_key", default=False)
    is_key = name in key or (k is True) or str(k).upper() in _TRUE
    ref = _get(fspec, "ref", "ref_field", "reffield", "reference_field")
    return FieldInfo(name=name, ddic_type=ddic, length=ln, decimals=dc, ref_field=str(ref).upper() if ref else None,
                     description=str(_get(fspec, "description", "text", "ddtext", "label", "scrtext_m", default="")),
                     is_key=is_key, source=Provenance.LIBRARY)


def _fms_from(section, d: ImportDraft, target: dict) -> None:
    items = section.items() if isinstance(section, dict) else [(None, x) for x in section or []]
    for name, spec in items:
        if not isinstance(spec, dict):
            continue
        fname = str(name or _get(spec, "name", "function", "fm", default="")).upper().strip()
        if _get(spec, "unknown") in (True, "true", "yes"):
            d.unknown.append(fname)
            continue
        kind = str(_get(spec, "kind", default="opaque")).lower()
        if kind not in FM_KINDS:
            d.warnings.append(f"{fname}: kind {kind!r} not one of {', '.join(FM_KINDS)}; set to opaque")
            kind = "opaque"
        tables = _get(spec, "tables", default=[])
        clean = {"kind": kind, "output": _get(spec, "output"), "sql": _get(spec, "sql"), "note": _get(spec, "note", "description"),
                 "tables": [str(t).upper() for t in (tables.split(",") if isinstance(tables, str) else tables or [])],
                 "table_param": _get(spec, "table_param"), "fieldcat_param": _get(spec, "fieldcat_param")}
        if not fname:
            continue
        target[fname] = fm_from_spec(fname, {k: v for k, v in clean.items() if v not in (None, "", [])}, Provenance.LIBRARY)


# ---------------------------------------------------------------------- zip (a whole pack or several files)
def _parse_zip(data: bytes, default_table: str | None, origin: str) -> ImportDraft:
    import zipfile

    d = ImportDraft(format="zip")
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        for info in z.infolist():
            n = info.filename
            if info.is_dir() or n.startswith("__MACOSX") or info.file_size > 20 * 1024 * 1024:
                continue
            low = n.lower()
            if low.endswith(("manifest.yaml", "history.jsonl")) or "/gold/" in f"/{low}":
                continue
            if low.endswith((".yaml", ".yml", ".json", ".csv", ".tsv", ".txt", ".md", ".abap", ".xlsx")):
                try:
                    sub = parse(z.read(info), n, default_table, origin)
                    d.merge(sub)
                except Exception as exc:  # noqa: BLE001
                    d.warnings.append(f"{n}: {exc}")
    if d.format == "mixed":
        d.format = "zip"
    return d


# ---------------------------------------------------------------------- preview
def preview(d: ImportDraft, kb) -> dict:
    """JSON for the UI: every parsed item with its status against the current knowledge (new / update / same)."""
    tables = []
    for name, t in sorted(d.tables.items()):
        cur = kb.table(name)
        if cur is None:
            status, added, changed = "new", len(t.fields), 0
        else:
            added = sum(1 for f in t.fields if f not in cur.fields)
            changed = sum(1 for f, x in t.fields.items() if f in cur.fields and (cur.fields[f].ddic_type, cur.fields[f].length) != (x.ddic_type, x.length))
            status = "update" if added or changed or (t.description and t.description != cur.description) else "same"
        tables.append({"name": name, "status": status, "current_source": cur.source.value if cur else None, "added_fields": added,
                       "changed_fields": changed, **table_json(t)})
    fms = []
    for section, items in (("function_modules", d.fms), ("methods", d.methods)):
        for name, f in sorted(items.items()):
            cur = kb.fm(name) if section == "function_modules" else kb.method(name)
            fms.append({"name": name, "section": section, "status": "new" if cur is None else "update", "current_source": cur.source.value if cur else None,
                        **f.model_dump(mode="json", exclude={"name"})})
    return {"format": d.format, "tables": tables, "fms": fms, "fm_sources": [{"name": n, "lines": s.count("\n") + 1} for n, s in d.fm_sources.items()],
            "s4": [m.model_dump(mode="json") for m in d.s4], "unknown": sorted(set(d.unknown)), "warnings": d.warnings[:200],
            "warning_count": len(d.warnings)}


def table_json(t: TableInfo) -> dict:
    return {"description": t.description, "module": t.module, "release": t.release, "key": t.key, "fields_complete": t.fields_complete,
            "text_table": t.text_table, "text_of": t.text_of, "source": t.source.value, "origin": t.origin, "updated": t.updated,
            "fields": [{"name": f.name, "ddic_type": f.ddic_type, "length": f.length, "decimals": f.decimals, "ref_field": f.ref_field,
                        "description": f.description, "is_key": f.is_key or f.name in t.key, "source": f.source.value} for f in t.fields.values()]}


def table_from_json(obj: dict, origin: str) -> TableInfo:
    """Inverse of table_json (the UI sends edited drafts back in this shape)."""
    d = ImportDraft()
    _tables_from([{**obj, "fields": obj.get("fields") or []}], d)
    if not d.tables:
        raise ValueError("; ".join(d.warnings) or "invalid table")
    t = next(iter(d.tables.values()))
    t.fields_complete = bool(obj.get("fields_complete"))
    t.origin = origin
    if not t.key:
        t.key = [n for n, f in t.fields.items() if f.is_key]
    return t
